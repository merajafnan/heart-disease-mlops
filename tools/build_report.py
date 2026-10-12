"""Build the 10-page assignment report from measured results and Windows evidence."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "reports" / "figures"
SHOTS = ROOT / "screenshots"
METADATA = ROOT / "artifacts" / "model_metadata.json"
SHOT_FILES = [
    "04-mlflow.png", "05-tests-ci.png", "05-tests-ci-artifacts.png",
    "06-docker.png", "07-docker-prediction.png",
    "08-kubernetes.png", "09-kubernetes-prediction.png", "10-monitoring.png",
]


def font(size=26, bold=False):
    name = "Arial Bold.ttf" if bold else "Arial.ttf"
    for candidate in (name, "/System/Library/Fonts/Supplemental/Arial.ttf"):
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default()


def architecture_image():
    path = FIGURES / "architecture.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (1350, 490), "white")
    draw = ImageDraw.Draw(canvas)
    boxes = [
        (30, 130, 245, 305, "UCI data\nEDA"),
        (300, 130, 515, 305, "Training\nTwo models"),
        (570, 130, 785, 305, "MLflow\nModel"),
        (840, 130, 1055, 305, "FastAPI\nDocker"),
        (1110, 130, 1325, 305, "Kubernetes\nMonitoring"),
    ]
    for x1, y1, x2, y2, label in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="#eaf1f7", outline="#526b80", width=3)
        bbox = draw.multiline_textbbox((0, 0), label, font=font(28, True), spacing=9, align="center")
        w, h = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.multiline_text((x1 + (x2-x1-w)/2, y1 + (y2-y1-h)/2), label,
                            fill="#172b3c", font=font(28, True), spacing=9, align="center")
    for x in (250, 520, 790, 1060):
        draw.line((x, 215, x+42, 215), fill="#526b80", width=5)
        draw.polygon(((x+42, 215), (x+28, 205), (x+28, 225)), fill="#526b80")
    draw.text((335, 55), "GitHub Actions: lint  |  tests  |  training  |  image build",
              fill="#172b3c", font=font(32, True))
    draw.text((890, 365), "Logs + /metrics + /monitor", fill="#172b3c", font=font(30))
    canvas.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, heading in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = heading
        set_cell_shading(cell, "263746")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    doc.add_paragraph()
    return table


def para(doc, text, style=None):
    return doc.add_paragraph(text, style=style)


def code_lines(doc, lines):
    block = doc.add_paragraph()
    block.paragraph_format.left_indent = Inches(0.18)
    block.paragraph_format.space_after = Pt(7)
    for index, line in enumerate(lines):
        run = block.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        if index < len(lines) - 1:
            run.add_break()


def deliverable(doc, number, title, details):
    item = doc.add_paragraph()
    item.paragraph_format.left_indent = Inches(0.22)
    item.paragraph_format.first_line_indent = Inches(-0.22)
    item.paragraph_format.space_after = Pt(1)
    item.add_run(f"{number}. {title}: ").bold = True
    item.add_run(details)


def heading(doc, title):
    doc.add_heading(title, level=1)


def figure(doc, filename, caption, width=5.7):
    path = FIGURES / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    caption_p = doc.add_paragraph(caption, style="Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def evidence(doc, filename, description, width=5.6, max_height=2.2):
    path = SHOTS / filename
    if path.is_file():
        with Image.open(path) as source:
            aspect = source.height / source.width
        image_width = min(width, max_height / aspect)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(image_width))
        c = doc.add_paragraph(description, style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph()
        p.style = "Caption"
        p.add_run(f"Windows evidence pending: {description} ({filename})")


def page(doc, title):
    doc.add_page_break()
    heading(doc, title)


def build(final=False, repo_url=""):
    if not METADATA.is_file():
        raise FileNotFoundError("Run python -m heart_mlops.train before building the report")
    if final and (not repo_url or any(not (SHOTS / name).is_file() for name in SHOT_FILES)):
        raise ValueError("Final report requires --repo-url and all Windows evidence screenshots")
    meta = json.loads(METADATA.read_text(encoding="utf-8"))
    architecture_image()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.67)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 11)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].font.bold = True
    styles["Caption"].font.name = "Arial"
    styles["Caption"].font.size = Pt(8)
    styles["Caption"].font.color.rgb = RGBColor(60, 60, 60)
    title_style = styles["Title"]._element
    for border in title_style.xpath(".//w:pBdr"):
        border.getparent().remove(border)

    doc.add_paragraph("Heart Disease Classification MLOps Experimental Learning Report", style="Title")
    para(doc, "This report covers model tests, experiment records, the API, and local Kubernetes deployment.")
    para(doc, "Prepared for the MLOps experimental learning assignment")
    para(doc, "Dataset: UCI processed Cleveland Heart Disease records")
    doc.add_heading("Project summary", level=1)
    para(doc, "This project classifies the UCI processed Cleveland heart disease records. It downloads and cleans the data, compares two models, and records the results in MLflow. FastAPI serves the selected model. Docker Desktop Kubernetes runs the API on the Windows laptop." + (" The report includes screenshots from that run." if final else " Add screenshots after the Windows run."))
    add_table(doc, ["Measure", "Windows run result" if final else "Development run result"], [
        ("Clean records", meta["data_summary"]["clean_rows"]),
        ("Selected model", meta["selected_model"].replace("_", " ").title()),
        ("Held out test accuracy", f"{meta['test_metrics']['accuracy']:.3f}"),
        ("Held out test ROC-AUC", f"{meta['test_metrics']['roc_auc']:.3f}"),
    ])
    para(doc, ("The Windows run used Python " if final else "The development run used Python ") + meta["python_version"] + ("." if final else " Confirm these values after the Windows run."))
    doc.add_heading("Deliverables and file locations", level=1)
    deliverable(doc, 1, "Code, Dockerfile, dependencies", "heart_mlops/ contains the code. Dockerfile builds the API image. requirements.txt and requirements-api.txt list the packages.")
    deliverable(doc, 2, "Cleaned data and download instructions", "data/processed/ has the clean CSV and data_summary.json. heart_mlops/data.py downloads and cleans the UCI file. RUNBOOK.md gives the command.")
    deliverable(doc, 3, "EDA, training, inference", "heart_mlops/ has eda.py, train.py, predict.py, and api.py. reports/figures/ has the plots. artifacts/model.joblib has the fitted model.")
    deliverable(doc, 4, "Unit tests", "tests/ has test_data.py, test_model.py, and test_api.py. These tests cover the data, model, and API.")
    deliverable(doc, 5, "GitHub Actions workflow", ".github/workflows/ci.yml runs Ruff, Pytest, data preparation, EDA, model training, and a Docker build. It uploads the run files.")
    deliverable(doc, 6, "Deployment manifests", "k8s/deployment.yaml starts the API pod. k8s/service.yaml gives it a local LoadBalancer Service on port 8000.")
    deliverable(doc, 7, "Screenshot folder", "screenshots/ has the ten main Windows screenshots and extra MLflow views. screenshots/README.md lists the main images.")
    deliverable(doc, 8, "Final written report", "reports/Final_Report.docx is this ten-page report. It includes the results, system diagram, and run screenshots.")
    deliverable(doc, 9, "Local API access instructions", "README.md and RUNBOOK.md give the localhost:8000 URLs and test commands. The API has no public URL.")

    page(doc, "1 Objective and system architecture")
    para(doc, "This task has two classes. Target 0 means no disease. Source targets 1 through 4 mean disease is present. The API takes 13 input values. It returns a class and a model probability.")
    figure(doc, "architecture.png", "Figure 1  Assignment system architecture", 6.6)
    para(doc, "The training script fits preprocessing and the classifier in one pipeline. MLflow stores the run results. The Docker image contains the selected pipeline. Docker Desktop Kubernetes runs one API pod. The service reports request counts and writes request logs.")
    doc.add_heading("Setup and local API access", level=2)
    para(doc, "The final run used Ubuntu 24.04 in WSL 2, Python 3.12.3, and Docker Desktop. From the repository root, run these commands in Ubuntu:")
    code_lines(doc, [
        "python3 -m venv ~/.venvs/heart-disease-mlops",
        "source ~/.venvs/heart-disease-mlops/bin/activate",
        "python -m pip install -r requirements.txt",
        "python -m heart_mlops.data",
        "python -m heart_mlops.eda",
        "ruff check .",
        "python -m pytest -q",
        "python -m heart_mlops.train",
    ])
    para(doc, "Build the API image in WSL. Then apply the two Kubernetes files from Windows PowerShell. RUNBOOK.md gives the commands and identifies the shell for each step.")
    code_lines(doc, [
        "docker build --build-arg PYTHON_VERSION=3.12 -t heart-disease-api:local .",
        "kubectl apply -f k8s\\deployment.yaml",
        "kubectl apply -f k8s\\service.yaml",
    ])
    para(doc, "The service is at http://localhost:8000 while Docker Desktop Kubernetes runs. Use /health to check readiness. Send sample_request.json to POST /predict. Open /monitor for counts or /metrics for Prometheus text.")

    page(doc, "2 Data source and preparation")
    para(doc, "The UCI processed Cleveland file has 303 rows, 13 inputs, and a target from 0 to 4. The download script checks the SHA-256 digest. It converts question marks to missing values. It removes incomplete or duplicate rows. It maps targets 1 through 4 to class 1.")
    add_table(doc, ["Data check", "Observed result"], [
        ("Raw rows", meta["data_summary"]["raw_rows"]),
        ("Rows after cleaning", meta["data_summary"]["clean_rows"]),
        ("Rows removed", meta["data_summary"]["removed_missing_or_duplicate_rows"]),
        ("Absent / present", f"{meta['data_summary']['negative_count']} / {meta['data_summary']['positive_count']}"),
    ])
    para(doc, "The repository has the clean CSV and a JSON file with row counts. The model pipeline can fill missing values in future requests. It learns these values from each training fold. The test rows do not set them.")
    para(doc, "Numeric inputs are age, resting blood pressure, cholesterol, maximum heart rate, and ST depression. The pipeline encodes the other eight inputs as categories. It does not use the target as an input.")

    page(doc, "3 Exploratory data analysis")
    para(doc, "The clean data has 160 class 0 rows and 137 class 1 rows. These plots show class counts and numeric values before model training.")
    figure(doc, "class_balance.png", "Figure 2  Cleaned dataset class balance", 4.0)
    figure(doc, "numeric_histograms.png", "Figure 3  Numeric predictor distributions", 5.2)

    page(doc, "4 Feature engineering and visual findings")
    para(doc, "The heatmap shows correlations between numeric inputs. It does not select features or show clinical cause. The box plots compare age and maximum heart rate by class. The small dataset limits these visual comparisons.")
    figure(doc, "correlation_heatmap.png", "Figure 4  Pairwise numeric correlations", 4.6)
    figure(doc, "feature_by_target.png", "Figure 5  Age and maximum heart rate by class", 5.0)

    page(doc, "5 Model design and selection")
    para(doc, "The script uses seed 42 to split the data. It puts 80% in training and 20% in testing. Five stratified folds test each parameter set. Numeric inputs use median imputation and scaling. Categorical inputs use mode imputation and one-hot encoding. The pipeline fits preprocessing inside each fold.")
    para(doc, "Logistic Regression is the linear baseline. Random Forest is the nonlinear model. The script selects the model with the highest mean validation ROC-AUC. It also records accuracy, precision, and recall.")
    rows = []
    for name, values in meta["candidate_cv_metrics"].items():
        rows.append((name.replace("_", " ").title(), *(f"{values[key]:.3f}" for key in ("accuracy", "precision", "recall", "roc_auc"))))
    add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "ROC-AUC"], rows)
    figure(doc, "random_forest_cv_scores.png", "Figure 6  Random Forest candidate validation scores", 5.8)

    page(doc, "6 Held out test evaluation")
    test = meta["test_metrics"]
    tn, fp = test["confusion_matrix"][0]
    fn, tp = test["confusion_matrix"][1]
    para(doc, "Cross validation selected " + meta["selected_model"].replace("_", " ") + f". The script then tested it once on {tn+fp+fn+tp} held-out rows. The confusion matrix shows {tn} true negatives, {fp} false positives, {fn} false negatives, and {tp} true positives.")
    add_table(doc, ["Metric", "Test result"], [
        ("Accuracy", f"{test['accuracy']:.3f}"), ("Precision", f"{test['precision']:.3f}"),
        ("Recall", f"{test['recall']:.3f}"), ("ROC-AUC", f"{test['roc_auc']:.3f}"),
    ])
    figure(doc, "test_confusion_matrix.png", "Figure 7  Test confusion matrix", 3.3)
    figure(doc, "test_roc_curve.png", "Figure 8  Test ROC curve", 3.3)

    page(doc, "7 Experiment tracking and reproducibility")
    para(doc, "MLflow stores one run for each candidate model and one run for the selected model. Candidate runs include parameters, validation scores, search results, and plots. The selected run includes test scores, figures, metadata, and the fitted model. SQLite stores run metadata. Artifact files remain on disk. GitHub Actions uploads both.")
    evidence(doc, "04-mlflow.png", "Figure 9  MLflow candidate and selected run metrics", 5.7, 3.6)
    para(doc, "The joblib file contains the fitted preprocessing and classifier. Model metadata lists the source checksum, selected parameters, inputs, seed, package version, and test scores. The requirements and download script support a new run.")

    page(doc, "8 Quality checks and continuous integration")
    para(doc, "Unit tests check target mapping, missing data, invalid targets, model output, API inputs, readiness, and monitoring endpoints. GitHub Actions runs Ruff, Pytest, data preparation, EDA, model training, and a Docker build. A failed required step marks the job as failed. The upload step still saves model and experiment files.")
    evidence(doc, "05-tests-ci.png", "Figure 10  Passing GitHub Actions workflow", 5.5)
    evidence(doc, "05-tests-ci-artifacts.png", "Figure 11  Uploaded workflow artifact", 5.1, 1.1)
    evidence(doc, "06-docker.png", "Figure 12  Local Docker image build", 5.0)
    evidence(doc, "07-docker-prediction.png", "Figure 13  Local container prediction and request logs", 5.4, 1.55)

    page(doc, "9 Deployment monitoring and limitations")
    para(doc, "Docker Desktop Kubernetes runs one API pod. The deployment sets health probes and CPU and memory limits. A LoadBalancer Service uses port 8000. The API has /predict, /health, /metrics, /monitor, and OpenAPI documentation. Logs contain method, route, status, and duration. They omit patient values.")
    para(doc, "The sample request returned class 1 from both the container and Kubernetes service. Its positive-class probability and confidence were 0.9837. This score is not a validated clinical risk estimate.")
    evidence(doc, "08-kubernetes.png", "Figure 14  Ready Kubernetes pod and LoadBalancer Service", 5.2)
    evidence(doc, "09-kubernetes-prediction.png", "Figure 15  Prediction from the deployed Kubernetes API", 5.3, 0.8)
    evidence(doc, "10-monitoring.png", "Figure 16  Live API monitoring dashboard", 5.2)
    para(doc, "The dataset is small and historical. This work has no external validation or fairness assessment. Model probabilities are not calibrated clinical risks. The Kubernetes API is available only on the local laptop.")
    para(doc, "Repository: " + (repo_url or "Add the GitHub URL after publishing the project."))
    para(doc, "Source: Janosi et al., Heart Disease, UCI Machine Learning Repository, 1988. https://doi.org/10.24432/C52P4X")

    target = ROOT / "reports" / ("Final_Report.docx" if final else "Report_Draft.docx")
    doc.save(target)
    print(target)
    return target


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--final", action="store_true")
    parser.add_argument("--repo-url", default="")
    args = parser.parse_args()
    build(final=args.final, repo_url=args.repo_url)
